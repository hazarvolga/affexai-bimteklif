from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path('/Users/hazarekiz/.codex/plugins/cache/openai-curated-remote/openai-templates/0.1.0/skills/artifact-template-system-design/assets/reference.docx')
OUTPUT = ROOT / 'docs/smartcheck/2026-07-23-live-bim-gap-report.docx'
doc = Document(str(REFERENCE))
body = doc._element.body
for child in list(body):
    if not child.tag.endswith('}sectPr'):
        body.remove(child)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(9.5)
styles['Normal'].paragraph_format.space_after = Pt(5)
for name, size, color in [('Title', 25, '123047'), ('Heading 1', 16, '123047'), ('Heading 2', 12, '00A6A6')]:
    styles[name].font.name = 'Aptos Display'
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

def para(text='', style=None, bold=False, color=None):
    item = doc.add_paragraph(style=style)
    run = item.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return item

def table(headers, rows):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        cell = result.rows[0].cells[i]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = result.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

para('TEKLİFPRO · LIVE BIM READINESS', bold=True, color='00A6A6')
title = para('Canlı Uygulama BIM GAP Raporu', style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
para('23 Temmuz 2026  |  localhost:3013  |  Playwright canlı tarayıcı kanıtı')
para('Karar: Uygulama demo ve kontrollü teknik pilot hazırlık seviyesinde; müşteri production pilotu için P0/P1 GAP’ler devam ediyor.', bold=True, color='123047')
para('1. Executive Summary', 'Heading 1')
para('Uygulama gerçek tarayıcıda başlatıldı. Public Academy, login, dashboard ve /app/smartcheck ekranları açıldı. SmartCheck ekranı evidence-first mesajını, Excel/CSV yükleme alanını, kolon eşlemesini ve kaynak hücre incelemesini görünür kılıyor.')
para('Canlı ekran kanıtı production pilot hazır olduğu anlamına gelmez. Gerçek müşteri dosyası, teknik reviewer kaydı, production S3 kanıtı ve browser upload-to-transfer kanıtı hâlâ gereklidir.')
para('2. Canlı Kanıt Günlüğü', 'Heading 1')
table(['Kanıt', 'Sonuç', 'Güven'], [('/academy', 'Açıldı; Academy navigasyonu ve CTA’lar görünür', 'Yüksek'), ('/login + admin oturumu', '/app dashboard’a yönlendi', 'Yüksek'), ('/app/smartcheck', 'Upload, mapping, kaynak kanıtı ve review alanı görünür', 'Yüksek'), ('HTTP + PostgreSQL harness', 'Gerçek endpoint zinciri ve tenant izolasyonu 1/1 geçti', 'Yüksek'), ('Console', 'favicon 404 ve development HMR websocket uyarısı', 'Yüksek')])
para('3. Current BIM Maturity', 'Heading 1')
table(['Boyut', 'Seviye', 'Bulgular'], [('People', 'Pilot hazırlığı', 'Reviewer veri sözleşmesi var; gerçek atama/onay yok'), ('Process', 'Orta', 'Upload → normalize → validate → review → transfer akışı var'), ('Information', 'Orta+', 'Source cell, raw/normalized değer, audit izi tutuluyor'), ('Governance', 'Orta-', 'Kural metadata var; canlı reviewer cockpit yok'), ('Technology', 'Orta+', 'Express/Vite/Drizzle/PostgreSQL/Worker ve storage adapter mevcut'), ('Evidence', 'Orta', 'Browser smoke ve HTTP harness var; müşteri dosyası yok')])
para('Bu değerlendirme ISO 19650 uyumluluğu veya sertifikasyonu iddiası değildir. EIR, BEP, CDE ve sözleşme bilgi gereksinimleri doğrulanmadan böyle bir sonuç verilemez.')
para('4. Information Requirements and Governance', 'Heading 1')
para('Her pilot dosyası müşteri/proje kimliği, izin ve retention, checksum, sheet/header mapping, source row/cell, raw/normalized değer, rule version, reviewer kararı ve aktarım idempotency sonucunu taşımalıdır.')
para('Önerilen kontrollü akış: Customer consent → upload quarantine → checksum/signature → parse → normalize → validate → human review → technical exception/rejection → approved transfer → quotation revision/PDF.')
para('5. Live UI / High-End Visual Review', 'Heading 1')
para('Seçilen arketip: Ethereal Glass + Asymmetrical Bento. SmartCheck alanının evidence-first konumlandırması güçlü; ancak demo denetimi ile gerçek import aynı sayfada daha sert ayrıştırılmalı, bekleyen feature gate ve reviewer/permission provenance paneli görünür olmalıdır.')
para('Canlı console’da favicon 404 gözlendi. Development HMR websocket uyarıları production smoke sinyalinden ayrıştırılmalı.')
para('6. Öncelikli GAP’ler', 'Heading 1')
table(['Öncelik', 'GAP', 'Çıkış kriteri'], [('P0', 'Gerçek MEP browser upload → transfer kanıtı', 'İzinli dosyada tam zincir ve PDF/regression kanıtı'), ('P0', 'Reviewer ve provenance görünürlüğü', 'İsim, scope, karar zamanı ve rule version görünür'), ('P0', 'Feature flag + allowlist canlı doğrulaması', 'Pilot şirketi dışında 404; pilot şirketinde akış açık'), ('P1', 'Sheet/header preview ve MEP sözlüğü', 'Beş dosyada mapping ve ≥%90 manuel yeniden girişsiz parse'), ('P1', 'Queue/operator/KPI görünürlüğü', 'Retry, bulgu, süre ve aktarım metrikleri kayıtlı'), ('P2', 'CDE/BEP/EIR şablon entegrasyonu', 'Proje sözleşmesiyle doğrulanmış bilgi gereksinimleri')])
para('7. Roadmap and Decisions', 'Heading 1')
para('Sıradaki adım: gerçek müşteri dosyası gelmeden önce P0 browser smoke ve provenance panelini tamamlamak; ardından beş dosyalı pilot admission başlatmak. Karar sahipleri: product owner, teknik reviewer, müşteri veri sahibi ve storage/DevOps sahibi.')
para('Confidence: Teknik route/harness bulguları yüksek; BIM saha genellenebilirliği orta-düşük, çünkü gerçek müşteri MEP örneklemi henüz yok.')
footer = doc.sections[0].footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.text = 'TeklifPro · Live BIM Readiness · Internal pilot evidence'
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(OUTPUT))
print(OUTPUT)
